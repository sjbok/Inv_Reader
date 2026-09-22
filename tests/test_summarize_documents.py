import json
import sys
import tempfile
import types
import unittest
from contextlib import redirect_stdout
from datetime import date
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Color, Font

from src.summarize_documents import (
    ConfidenceAnalysis,
    DocumentContent,
    DATABASE_HEADERS,
    FileProcessingStatus,
    OllamaClient,
    OllamaRuntime,
    OUTPUT_WORKBOOK_NAME,
    Item,
    ProcessingReport,
    PurchaseOrder,
    bundled_ollama_path,
    configured_model,
    confidence_level,
    extract_purchase_order,
    iter_documents,
    main,
    output_name,
    process_documents,
    read_document,
    show_issue_popups,
    summarize_content,
    _parse_order_date,
)


class FakeClient:
    def __init__(self):
        self.prompts = []

    def chat(self, prompt, images=None):
        self.prompts.append((prompt, images))
        return "Summary {}".format(len(self.prompts))


class FakeResponse:
    def __init__(self, body):
        self.body = body

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def read(self):
        return json.dumps(self.body).encode("utf-8")


class SummarizerTests(unittest.TestCase):
    def _write_order(self, path, memo="배송 전 담당자에게 연락 바랍니다.",
                     order_date="2026-09-03",
                     include_example=False, example_color="FF7F7F7F"):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "발주서"
        rows = [
            ("발주일자", order_date),
            ("발주처", "테스트 업체"),
            ("수령인", "홍길동", "", "수령인 연락처", "010-0000-0000"),
            ("배송지 주소", "서울시 테스트구 테스트로 1"),
            ("품목코드", "품목명[규격]", "수량", "단위"),
            ("A-1", "테스트 상품 [중형]", 2, "개"),
            ("A-2", "두 번째 상품", 3, "박스"),
            ("특기사항 • 배송 메모", memo),
        ]
        if include_example:
            rows.insert(5, ("EXAMPLE", "예시 상품", 99, "개"))
        for row in rows:
            worksheet.append(row)
        if include_example:
            for cell in worksheet[6]:
                if cell.value is not None:
                    cell.font = Font(italic=True, color=example_color)
        workbook.save(path)
        workbook.close()

    def test_xlsx_orders_are_appended_to_daily_database_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            source = input_dir / "order.xlsx"
            self._write_order(source)

            client = FakeClient()
            self.assertEqual(process_documents(input_dir, output_dir, client), 0)

            database = output_dir / OUTPUT_WORKBOOK_NAME
            workbook = load_workbook(database, data_only=True)
            self.assertEqual(workbook.sheetnames, ["09.03"])
            self.assertEqual(list(workbook["09.03"].values), [
                DATABASE_HEADERS,
                ("홍길동", "010-0000-0000", "서울시 테스트구 테스트로 1",
                 "(A-1) 테스트 상품 [중형] 2개", "테스트 업체", "배송 전 담당자에게 연락 바랍니다."),
                ("홍길동", "010-0000-0000", "서울시 테스트구 테스트로 1",
                 "(A-2) 두 번째 상품 3박스", "테스트 업체", "배송 전 담당자에게 연락 바랍니다."),
            ])
            workbook.close()
            self.assertEqual(client.prompts, [])

    def test_existing_database_is_reused_when_processing_another_order(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            output_dir.mkdir()
            first = input_dir / "first.xlsx"
            self._write_order(first)
            self.assertEqual(process_documents(input_dir, output_dir), 0)

            second_input = root / "second-input"
            second_input.mkdir()
            second = second_input / "second.xlsx"
            self._write_order(second, memo="두 번째 주문")
            self.assertEqual(process_documents(second_input, output_dir), 0)

            workbook = load_workbook(output_dir / OUTPUT_WORKBOOK_NAME, data_only=True)
            rows = list(workbook["09.03"].values)
            self.assertEqual(len(rows), 5)
            self.assertEqual(rows[-1][-1], "두 번째 주문")
            workbook.close()

    def test_order_dates_are_normalized_from_common_formats(self):
        self.assertEqual(_parse_order_date("26/09/02"), date(2026, 9, 2))
        self.assertEqual(_parse_order_date("26/9/3"), date(2026, 9, 3))
        self.assertEqual(_parse_order_date("2026/9/3"), date(2026, 9, 3))
        self.assertEqual(_parse_order_date("02/09/2026"), date(2026, 9, 2))
        self.assertEqual(_parse_order_date("12/31/26"), date(2026, 12, 31))

    def test_xlsx_text_normalizes_order_date_for_model_reading(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "order.xlsx"
            self._write_order(path, order_date="26/9/3")

            content = read_document(path)

            self.assertIn("발주일자\t2026-09-03", content.text)

    def test_input_filenames_are_logged_by_order_date_and_duplicates_are_skipped(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            source = input_dir / "order.xlsx"
            self._write_order(source, order_date="26/9/3")

            self.assertEqual(process_documents(input_dir, output_dir), 0)
            log = output_dir / "2026-09-03.txt"
            self.assertEqual(log.read_text(encoding="utf-8"), "order.xlsx\n")

            output = StringIO()
            report = ProcessingReport([], [])
            with redirect_stdout(output):
                self.assertEqual(process_documents(input_dir, output_dir, report=report), 0)
            self.assertIn("DUPLICATE", output.getvalue())
            self.assertIn("order.xlsx", output.getvalue())
            self.assertEqual(report.duplicates, ["order.xlsx"])
            self.assertEqual(log.read_text(encoding="utf-8"), "order.xlsx\n")

            workbook = load_workbook(output_dir / OUTPUT_WORKBOOK_NAME, data_only=True)
            self.assertEqual(workbook["09.03"].max_row, 3)
            workbook.close()

    def test_low_confidence_xlsx_is_reported_for_human_review(self):
        order = PurchaseOrder(
            order_date=None,
            recipient="홍길동",
            phone="010-0000-0000",
            address="서울시",
            company="업체",
            memo="메모",
            items=[Item(code="A-1", name="상품", quantity="1", unit="개")],
        )
        self.assertEqual(confidence_level(order), 87.5)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            source = input_dir / "needs-review.xlsx"
            self._write_order(source, memo="")

            output = StringIO()
            report = ProcessingReport([], [])
            with redirect_stdout(output):
                self.assertEqual(process_documents(input_dir, output_dir, report=report), 0)
            self.assertIn("needs-review.xlsx", output.getvalue())
            self.assertIn("HUMAN REVIEW REQUIRED", output.getvalue())
            self.assertEqual([(result.filename, result.confidence) for result in report.low_confidence], [
                ("needs-review.xlsx", 88.89),
            ])

    def test_file_status_callback_reports_reading_and_low_confidence(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            self._write_order(input_dir / "needs-review.xlsx", memo="")

            statuses = []
            self.assertEqual(process_documents(
                input_dir,
                output_dir,
                on_file_status=statuses.append,
            ), 0)

            self.assertEqual([(status.filename, status.status) for status in statuses], [
                ("needs-review.xlsx", "reading"),
                ("needs-review.xlsx", "low_confidence"),
            ])
            self.assertEqual(statuses[-1].confidence, 88.89)

    def test_file_status_callback_marks_duplicate_after_reading(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            self._write_order(input_dir / "order.xlsx")
            self.assertEqual(process_documents(input_dir, output_dir), 0)

            statuses = []
            self.assertEqual(process_documents(
                input_dir,
                output_dir,
                on_file_status=statuses.append,
            ), 0)

            self.assertEqual(statuses[-1], FileProcessingStatus("order.xlsx", "duplicate"))

    def test_structured_xlsx_extraction_ignores_unused_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "order.xlsx"
            self._write_order(path)
            order = extract_purchase_order(path)

            self.assertEqual(order.company, "테스트 업체")
            self.assertEqual(len(order.items), 2)
            self.assertEqual(order.items[0].code, "A-1")

    def test_structured_xlsx_extraction_ignores_grey_italic_example_row(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "order.xlsx"
            self._write_order(path, include_example=True)

            order = extract_purchase_order(path)

            self.assertEqual([item.code for item in order.items], ["A-1", "A-2"])

    def test_structured_xlsx_extraction_ignores_indexed_grey_italic_example_row(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "order.xlsx"
            self._write_order(path, include_example=True, example_color=Color(indexed=13))

            order = extract_purchase_order(path)

            self.assertEqual([item.code for item in order.items], ["A-1", "A-2"])

    def test_summarizes_text_into_flat_output_file(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            (input_dir / "nested").mkdir()
            source = input_dir / "nested" / "invoice.txt"
            source.write_text("Invoice total: $42", encoding="utf-8")

            client = FakeClient()
            failures = process_documents(input_dir, output_dir, client)

            self.assertEqual(failures, 0)
            self.assertEqual((output_dir / "nested__invoice.txt").read_text(encoding="utf-8"), "Summary 1\n")
            self.assertIn("Invoice total: $42", client.prompts[0][0])

    def test_long_text_is_summarized_in_parts_then_combined(self):
        client = FakeClient()
        summary = summarize_content(DocumentContent(text="first\n\nsecond\n\nthird"), client, chunk_size=8)

        self.assertEqual(summary, "Summary 4")
        self.assertEqual(len(client.prompts), 4)
        self.assertIn("part 1/3", client.prompts[0][0])
        self.assertIn("PART 3 SUMMARY", client.prompts[-1][0])

    def test_image_content_is_sent_to_the_model(self):
        client = FakeClient()
        summarize_content(DocumentContent(images=["encoded-image"]), client)

        self.assertEqual(client.prompts[0][1], ["encoded-image"])

    def test_text_requests_disable_broken_qwen3_vl_thinking_template(self):
        response = FakeResponse({"message": {"content": "Summary"}})
        client = OllamaClient()

        with patch("src.summarize_documents.urllib.request.urlopen", return_value=response) as urlopen:
            self.assertEqual(client.chat("Summarize this"), "Summary")

        payload = json.loads(urlopen.call_args.args[0].data.decode("utf-8"))
        self.assertTrue(payload["raw"])
        self.assertFalse(payload["think"])
        self.assertEqual(payload["messages"][-1], {
            "role": "assistant",
            "content": "<think>\n\n</think>\n\n",
        })
        self.assertGreater(payload["options"]["num_predict"], 0)

    def test_portable_runtime_starts_and_stops_its_own_server(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            executable = root / "ollama" / "ollama.exe"
            executable.parent.mkdir()
            executable.write_bytes(b"fake")
            models = root / "models"
            process = types.SimpleNamespace(
                poll=lambda: None,
                terminate=lambda: None,
                wait=lambda timeout=None: 0,
            )
            response = FakeResponse({"models": [{"name": "qwen3-vl:8b"}]})
            with patch("src.summarize_documents.subprocess.Popen", return_value=process) as popen:
                with patch("src.summarize_documents.urllib.request.urlopen", return_value=response):
                    runtime = OllamaRuntime(
                        executable, models, "http://127.0.0.1:11435", "qwen3-vl:8b",
                        root / "run.log", startup_timeout=1,
                    )
                    runtime.start()
                    runtime.stop()

            self.assertEqual(popen.call_args.args[0], [str(executable), "serve"])
            self.assertEqual(popen.call_args.kwargs["env"]["OLLAMA_HOST"], "127.0.0.1:11435")
            self.assertEqual(popen.call_args.kwargs["env"]["OLLAMA_MODELS"], str(models))
            self.assertTrue(models.is_dir())

    def test_portable_runtime_pulls_a_missing_model_in_the_background(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            executable = root / "ollama.exe"
            executable.write_bytes(b"fake")
            serve = types.SimpleNamespace(
                poll=lambda: None,
                terminate=lambda: None,
                wait=lambda timeout=None: 0,
            )
            pull = types.SimpleNamespace(wait=lambda: 0)
            responses = [
                FakeResponse({"models": []}),
                FakeResponse({"models": []}),
                FakeResponse({"models": [{"name": "qwen3-vl:8b"}]}),
            ]
            with patch("src.summarize_documents.subprocess.Popen",
                       side_effect=[serve, pull]) as popen:
                with patch("src.summarize_documents.urllib.request.urlopen",
                           side_effect=responses):
                    runtime = OllamaRuntime(
                        executable, root / "models", "http://127.0.0.1:11435", "qwen3-vl:8b",
                        root / "run.log", startup_timeout=1,
                    )
                    runtime.start()
                    runtime.stop()

            self.assertEqual(popen.call_args_list[1].args[0], [str(executable), "pull", "qwen3-vl:8b"])

    def test_portable_runtime_stops_the_windows_process_tree(self):
        process = types.SimpleNamespace(pid=321)
        runtime = OllamaRuntime(
            Path("ollama.exe"), Path("models"), "http://127.0.0.1:11435", "qwen3-vl:8b",
            Path("run.log"),
        )
        runtime.process = process

        with patch("src.summarize_documents.os.name", "nt"):
            with patch("src.summarize_documents.subprocess.run") as taskkill:
                runtime.stop()

        self.assertEqual(taskkill.call_args.args[0], ["taskkill", "/PID", "321", "/T", "/F"])

    def test_portable_runtime_can_find_ollama_beside_executable(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            nested = root / "ollama" / "ollama.exe"
            nested.parent.mkdir()
            nested.write_bytes(b"fake")

            self.assertEqual(bundled_ollama_path(root), nested)

    def test_portable_model_configuration_is_used_when_present(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "portable_config.json").write_text(
                '{"model": "custom-model:latest"}', encoding="utf-8"
            )

            self.assertEqual(configured_model(root), "custom-model:latest")

    def test_main_creates_runtime_folders_before_processing(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = StringIO()
            with patch("src.summarize_documents.application_root", return_value=root):
                with redirect_stdout(output):
                    self.assertEqual(main(["--no-ollama-management"]), 1)

            self.assertTrue((root / "input").is_dir())
            self.assertTrue((root / "output").is_dir())

    def test_windows_build_shows_one_popup_for_each_issue_category(self):
        report = ProcessingReport(
            ["old-order.xlsx", "nested/another-order.xlsx"],
            [ConfidenceAnalysis("uncertain.xlsx", 75.0)],
        )
        message_box = types.SimpleNamespace(MessageBoxW=lambda *args: None)
        fake_ctypes = types.SimpleNamespace(windll=types.SimpleNamespace(user32=message_box))

        with patch("src.summarize_documents.ctypes", fake_ctypes):
            with patch("src.summarize_documents.os.name", "nt"):
                with patch("src.summarize_documents.sys.frozen", True, create=True):
                    with patch.object(message_box, "MessageBoxW") as show_message:
                        show_issue_popups(report)

        self.assertEqual(show_message.call_count, 2)
        self.assertIn("old-order.xlsx", show_message.call_args_list[0].args[1])
        self.assertIn("uncertain.xlsx: 75.00%", show_message.call_args_list[1].args[1])

    def test_pdf_prefers_layout_aware_text_extraction(self):
        class FakePage:
            def __init__(self):
                self.sort = None

            def get_text(self, kind, sort=False):
                self.sort = sort
                return "layout text"

        class FakeDocument:
            def __init__(self, page):
                self.page = page

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def __iter__(self):
                return iter([self.page])

        page = FakePage()
        fake_fitz = types.SimpleNamespace(open=lambda path: FakeDocument(page))
        fake_pypdf = types.SimpleNamespace(
            PdfReader=lambda path: types.SimpleNamespace(
                pages=[types.SimpleNamespace(extract_text=lambda: "pypdf text")]
            )
        )

        with patch.dict(sys.modules, {"fitz": fake_fitz, "pypdf": fake_pypdf}):
            content = read_document(Path("invoice.pdf"))

        self.assertEqual(content.text, "[Page 1]\nlayout text")
        self.assertTrue(page.sort)

    def test_xlsx_extracts_nonempty_rows_and_preserves_sheet_names(self):
        class FakeWorksheet:
            title = "발주서"

            def iter_rows(self, values_only=False):
                self.values_only = values_only
                return iter([
                    ("발주일자", "2026-09-02", None),
                    (None, None, None),
                    ("품목", None, "수량"),
                    ("AED-200", None, 5),
                ])

        class FakeWorkbook:
            worksheets = [FakeWorksheet()]

            def close(self):
                self.closed = True

        workbook = FakeWorkbook()
        fake_openpyxl = types.SimpleNamespace(
            load_workbook=lambda path, read_only, data_only: workbook
        )
        with patch.dict(sys.modules, {"openpyxl": fake_openpyxl}):
            content = read_document(Path("purchase.xlsx"))

        self.assertEqual(content.text, (
            "[Sheet: 발주서]\n"
            "발주일자\t2026-09-02\n"
            "품목\t\t수량\n"
            "AED-200\t\t5"
        ))

    def test_prompt_instructs_model_to_read_english_and_korean(self):
        client = FakeClient()
        summarize_content(DocumentContent(text="English text\n\n한국어 문서"), client)

        self.assertIn("English, Korean, or a mixture of both", client.prompts[0][0])
        self.assertIn("한국어", client.prompts[0][0])

    def test_prompt_preserves_original_wording_and_prioritizes_key_fields(self):
        client = FakeClient()
        summarize_content(DocumentContent(text="이름: 홍길동\n업체명: Example Co."), client)

        prompt = client.prompts[0][0]
        for field in ("이름", "전화", "주소", "품목 (수량)", "업체명"):
            self.assertIn(field, prompt)
        self.assertIn("copy the wording exactly", prompt)
        self.assertIn("do not translate, romanize, or normalize it", prompt)

    def test_unsupported_files_are_not_discovered(self):
        with tempfile.TemporaryDirectory() as directory:
            input_dir = Path(directory) / "input"
            input_dir.mkdir()
            supported = input_dir / "notes.md"
            spreadsheet = input_dir / "purchase.xlsx"
            unsupported = input_dir / "archive.zip"
            supported.write_text("notes", encoding="utf-8")
            spreadsheet.write_bytes(b"xlsx")
            unsupported.write_bytes(b"zip")

            self.assertEqual(iter_documents(input_dir), [supported, spreadsheet])
            self.assertEqual(output_name(supported, input_dir), "notes.txt")

    def test_flattened_names_do_not_collide_with_separator_in_filename(self):
        with tempfile.TemporaryDirectory() as directory:
            input_dir = Path(directory) / "input"
            input_dir.mkdir()
            nested_dir = input_dir / "a"
            nested_dir.mkdir()
            nested = nested_dir / "b.txt"
            flat = input_dir / "a__b.txt"

            self.assertNotEqual(output_name(nested, input_dir), output_name(flat, input_dir))


if __name__ == "__main__":
    unittest.main()
